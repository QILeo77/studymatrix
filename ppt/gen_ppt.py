"""StudyMatrix 项目演示 PPT 生成脚本（python-pptx）
翡翠绿主题（贴合项目 UI #10B981），7 页初稿
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

# ---------- 主题色 ----------
EMERALD = RGBColor(0x10, 0xB9, 0x81)      # 主色：翡翠绿
EMERALD_DARK = RGBColor(0x05, 0x96, 0x69)  # 深绿
DARK = RGBColor(0x1F, 0x29, 0x37)          # 深灰文字
GRAY = RGBColor(0x6B, 0x72, 0x80)          # 次要文字
LIGHT = RGBColor(0xF9, 0xFA, 0xFB)         # 极浅灰背景
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE5, 0xE7, 0xEB)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "StudyMatrix项目演示.pptx"


def add_text(slide, text, left, top, width, height, size=16, bold=False,
             color=DARK, align=PP_ALIGN.LEFT, font="Microsoft YaHei", anchor=MSO_ANCHOR.TOP):
    """通用文本框"""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
        p.space_after = Pt(4)
    return tb


def add_rect(slide, left, top, width, height, fill=WHITE, line=None, radius=None):
    """圆角矩形卡片"""
    shape = slide.shapes.add_shape(
        5,  # MSO_SHAPE.ROUNDED_RECTANGLE
        left, top, width, height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    if radius is not None:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    shape.shadow.inherit = False
    return shape


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    return slide


def page_header(slide, title, subtitle=None):
    """内页标题：绿色小竖条 + 标题 + 副标题"""
    bar = add_rect(slide, Inches(0.6), Inches(0.55), Inches(0.09), Inches(0.52), fill=EMERALD)
    add_text(slide, title, Inches(0.85), Inches(0.42), Inches(9), Inches(0.6), size=26, bold=True, color=DARK)
    if subtitle:
        add_text(slide, subtitle, Inches(0.87), Inches(1.02), Inches(10), Inches(0.4), size=13, color=GRAY)
    return slide


def build_cover(prs):
    s = new_slide(prs)
    # 大面积白底 + 顶部绿色条
    add_rect(s, Emu(0), Emu(0), SLIDE_W, Inches(0.18), fill=EMERALD)
    # 左侧装饰：绿色圆点阵
    for i in range(5):
        for j in range(7):
            add_rect(s, Inches(0.8 + i * 0.32), Inches(2.2 + j * 0.32),
                     Inches(0.1), Inches(0.1), fill=RGBColor(0xD1, 0xFA, 0xE5))
    # 中央标题
    add_text(s, "智学矩阵 StudyMatrix", Inches(3.2), Inches(2.35), Inches(9.5), Inches(1.1),
             size=48, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    add_text(s, "多 Agent 智能复习规划助手 · 项目演示初稿", Inches(3.2), Inches(3.55), Inches(9.5), Inches(0.6),
             size=20, color=EMERALD_DARK, align=PP_ALIGN.CENTER)
    # 分隔小点
    add_rect(s, Inches(6.28), Inches(4.35), Inches(0.78), Inches(0.05), fill=EMERALD)
    add_text(s, "Flask + DeepSeek + SSE 流式推送 + 原生前端（Tailwind + marked.js）",
             Inches(2.2), Inches(4.6), Inches(11), Inches(0.5),
             size=14, color=GRAY, align=PP_ALIGN.CENTER)
    add_text(s, "2026 年 10 月", Inches(2.2), Inches(6.6), Inches(11), Inches(0.4),
             size=12, color=GRAY, align=PP_ALIGN.CENTER)


def build_content(prs, data):
    s = new_slide(prs)
    page_header(s, data["title"], data.get("subtitle"))
    items = data["items"]
    top = Inches(1.75)
    for item in items:
        # 浅灰卡片行
        add_rect(s, Inches(0.85), top, Inches(11.6), Inches(0.86), fill=LIGHT, line=BORDER, radius=0.12)
        # 绿色圆点
        add_rect(s, Inches(1.15), top + Inches(0.31), Inches(0.16), Inches(0.16), fill=EMERALD)
        add_text(s, item, Inches(1.55), top + Inches(0.12), Inches(10.6), Inches(0.66),
                 size=15, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
        top = top + Inches(1.06)


def build_data(prs, data):
    s = new_slide(prs)
    page_header(s, data["title"])
    # 三个指标卡
    metrics = data["metrics"]
    w = Inches(3.6)
    gap = Inches(0.45)
    left = Inches(0.85)
    for m in metrics:
        add_rect(s, left, Inches(1.8), w, Inches(1.9), fill=WHITE, line=BORDER, radius=0.1)
        add_text(s, m["label"], left, Inches(2.05), w, Inches(0.5),
                 size=14, color=GRAY, align=PP_ALIGN.CENTER)
        add_text(s, m["value"], left, Inches(2.55), w, Inches(0.9),
                 size=34, bold=True, color=EMERALD_DARK, align=PP_ALIGN.CENTER)
        left = left + w + gap
    # 下方要点
    top = Inches(4.15)
    for item in data["items"]:
        add_rect(s, Inches(0.85), top, Inches(0.14), Inches(0.14), fill=EMERALD)
        add_text(s, item, Inches(1.2), top - Inches(0.08), Inches(11), Inches(0.6), size=15, color=DARK)
        top = top + Inches(0.78)


def build_timeline(prs, data):
    s = new_slide(prs)
    page_header(s, data["title"], "SSE 流式推送，每一步前端实时可见")
    steps = data["steps"]
    n = len(steps)
    # 横向时间轴主线
    line_y = Inches(2.9)
    add_rect(s, Inches(1.0), line_y, Inches(11.4), Inches(0.05), fill=BORDER)
    seg = Inches(11.4 / n)
    for i, st in enumerate(steps):
        cx = Inches(1.0 + seg * i + seg / 2)
        # 圆形节点
        node = s.shapes.add_shape(9, cx - Inches(0.22), line_y - Inches(0.19), Inches(0.44), Inches(0.44))  # OVAL
        node.fill.solid()
        node.fill.fore_color.rgb = EMERALD
        node.line.fill.background()
        node.shadow.inherit = False
        # 节点序号
        tf = node.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(i + 1)
        r.font.size = Pt(16); r.font.bold = True
        r.font.color.rgb = WHITE; r.font.name = "Microsoft YaHei"
        # 标题（节点上方）
        add_text(s, st["title"], cx - seg / 2, line_y - Inches(1.05), seg, Inches(0.5),
                 size=16, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        # 描述（节点下方）
        add_text(s, st["desc"], cx - seg / 2 + Inches(0.1), line_y + Inches(0.45), seg - Inches(0.2), Inches(0.8),
                 size=13, color=GRAY, align=PP_ALIGN.CENTER)
    # 底部说明条
    add_rect(s, Inches(0.85), Inches(5.6), Inches(11.6), Inches(1.0), fill=LIGHT, line=BORDER, radius=0.1)
    add_text(s, "示例：发送「帮我制定复习计划」→ 灰字 [Router Agent: 正在识别意图...] → [Tool: 正在读取本地教务系统 schedule.json...] → 三列表格（时间/任务/目标）流式呈现",
             Inches(1.15), Inches(5.75), Inches(11.0), Inches(0.7), size=13, color=GRAY)


def build_split(prs, data):
    """移动端页：左侧要点 + 右侧手机线框"""
    s = new_slide(prs)
    page_header(s, data["title"], data.get("subtitle"))
    top = Inches(1.9)
    for item in data["items"]:
        add_rect(s, Inches(0.85), top, Inches(6.6), Inches(0.8), fill=LIGHT, line=BORDER, radius=0.12)
        add_rect(s, Inches(1.1), top + Inches(0.3), Inches(0.14), Inches(0.14), fill=EMERALD)
        add_text(s, item, Inches(1.45), top + Inches(0.08), Inches(5.9), Inches(0.64),
                 size=14, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
        top = top + Inches(0.98)

    # 右侧：手机线框（圆角矩形 + 顶栏 + 气泡）
    px, py, pw, ph = Inches(8.4), Inches(1.7), Inches(3.2), Inches(5.4)
    add_rect(s, px, py, pw, ph, fill=WHITE, line=DARK, radius=0.08)
    # 手机顶栏（绿色）
    add_rect(s, px + Inches(0.12), py + Inches(0.12), pw - Inches(0.24), Inches(0.5), fill=EMERALD, radius=0.18)
    add_text(s, "智学矩阵 · 待命中 ●", px + Inches(0.12), py + Inches(0.2), pw - Inches(0.24), Inches(0.35),
             size=10, color=WHITE, align=PP_ALIGN.CENTER)
    # AI 气泡（左）
    add_rect(s, px + Inches(0.25), py + Inches(0.9), Inches(1.9), Inches(0.75), fill=LIGHT, line=BORDER, radius=0.3)
    add_text(s, "AI: 帮你排好复习计划啦", px + Inches(0.35), py + Inches(1.02), Inches(1.7), Inches(0.5), size=9, color=DARK)
    # 用户气泡（右，绿）
    add_rect(s, px + Inches(1.05), py + Inches(1.85), Inches(1.9), Inches(0.6), fill=EMERALD, radius=0.3)
    add_text(s, "帮我制定复习计划", px + Inches(1.15), py + Inches(1.95), Inches(1.7), Inches(0.4), size=9, color=WHITE)
    # AI 回复表格气泡
    add_rect(s, px + Inches(0.25), py + Inches(2.65), Inches(2.6), Inches(1.5), fill=LIGHT, line=BORDER, radius=0.15)
    add_text(s, "| 时间 | 任务 | 目标 |\n| 19:00 | 高数 | 及格 |\n| 20:00 | 英语 | 六级 |",
             px + Inches(0.4), py + Inches(2.8), Inches(2.3), Inches(1.2), size=8.5, color=DARK)
    # 输入框 + 发送按钮
    add_rect(s, px + Inches(0.2), py + Inches(4.5), pw - Inches(0.4), Inches(0.55), fill=WHITE, line=BORDER, radius=0.5)
    add_text(s, "给智学矩阵发送消息…", px + Inches(0.45), py + Inches(4.62), Inches(1.8), Inches(0.35), size=9, color=GRAY)
    send = s.shapes.add_shape(9, px + pw - Inches(0.72), py + Inches(4.53), Inches(0.48), Inches(0.48))
    send.fill.solid(); send.fill.fore_color.rgb = EMERALD
    send.line.fill.background(); send.shadow.inherit = False
    add_text(s, "↑", px + pw - Inches(0.72), py + Inches(4.56), Inches(0.48), Inches(0.42),
             size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def build_comparison(prs, data):
    s = new_slide(prs)
    page_header(s, data["title"])
    lw = Inches(5.6)
    # 左：传统（灰）
    add_rect(s, Inches(0.85), Inches(1.8), lw, Inches(4.6), fill=LIGHT, line=BORDER, radius=0.08)
    add_text(s, data["left"]["title"], Inches(0.85), Inches(2.0), lw, Inches(0.5),
             size=18, bold=True, color=GRAY, align=PP_ALIGN.CENTER)
    top = Inches(2.75)
    for it in data["left"]["items"]:
        add_rect(s, Inches(1.25), top + Inches(0.1), Inches(0.13), Inches(0.13), fill=GRAY)
        add_text(s, it, Inches(1.6), top, Inches(4.4), Inches(0.6), size=15, color=GRAY)
        top = top + Inches(0.85)
    # 右：智学矩阵（绿）
    rx = Inches(6.9)
    add_rect(s, rx, Inches(1.8), lw, Inches(4.6), fill=RGBColor(0xEC, 0xFD, 0xF5), line=EMERALD, radius=0.08)
    add_text(s, data["right"]["title"], rx, Inches(2.0), lw, Inches(0.5),
             size=18, bold=True, color=EMERALD_DARK, align=PP_ALIGN.CENTER)
    top = Inches(2.75)
    for it in data["right"]["items"]:
        add_rect(s, rx + Inches(0.4), top + Inches(0.1), Inches(0.13), Inches(0.13), fill=EMERALD)
        add_text(s, it, rx + Inches(0.75), top, Inches(4.4), Inches(0.6), size=15, color=DARK)
        top = top + Inches(0.85)


def build_end(prs, data):
    s = new_slide(prs)
    add_rect(s, Emu(0), Emu(0), SLIDE_W, Inches(0.18), fill=EMERALD)
    add_text(s, data["title"], Inches(1), Inches(2.6), Inches(11.3), Inches(1.0),
             size=44, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    add_text(s, data["subtitle"], Inches(1), Inches(3.8), Inches(11.3), Inches(0.6),
             size=18, color=EMERALD_DARK, align=PP_ALIGN.CENTER)
    add_text(s, data.get("footer", ""), Inches(1), Inches(6.5), Inches(11.3), Inches(0.5),
             size=13, color=GRAY, align=PP_ALIGN.CENTER)


def main():
    cfg = json.loads((BASE / "ppt" / "content.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    for slide_cfg in cfg["slides"]:
        layout = slide_cfg["layout"]
        data = slide_cfg.get("data", {})
        if layout == "cover":
            build_cover(prs)
        elif layout == "content":
            build_content(prs, data)
        elif layout == "data":
            build_data(prs, data)
        elif layout == "timeline":
            build_timeline(prs, data)
        elif layout == "split":
            build_split(prs, data)
        elif layout == "comparison":
            build_comparison(prs, data)
        elif layout == "end":
            build_end(prs, data)

    prs.save(str(OUT))
    print("SAVED:", OUT)


if __name__ == "__main__":
    main()
